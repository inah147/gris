"""Testes da exportação do extrato em CSV e XLSX.

Cobrem o formato dos arquivos (pt-BR no CSV, tipos nativos no XLSX), a
proteção contra fórmulas vindas das descrições e o contrato do endpoint, que
exporta todas as linhas do filtro da tela, não só o lote carregado.
"""

import csv
import io
import json
from datetime import date, datetime

import frappe
from frappe.tests.utils import FrappeTestCase
from openpyxl import load_workbook

from gris.api.financeiro.transactions import (
	EXTRATO_MAX_PAGE_SIZE,
	exportar_extrato,
	gerar_csv_extrato,
	gerar_xlsx_extrato,
	get_extrato_colunas,
)

DOCTYPE = "Transacao Extrato Geral"

COLUNAS = [
	{"key": "descricao_reduzida", "label": "Descrição reduzida", "tipo": "texto"},
	{"key": "valor", "label": "Valor", "tipo": "moeda"},
	{"key": "timestamp_transacao", "label": "Data/Hora", "tipo": "datahora"},
	{"key": "data_transacao", "label": "Data da transação", "tipo": "data"},
	{"key": "transacao_revisada", "label": "Revisão", "tipo": "revisao"},
]

TRANSACAO = {
	"name": "TX-0001",
	"descricao_reduzida": '=HYPERLINK("http://x","clique")',
	"valor": -1234.5,
	"timestamp_transacao": datetime(2026, 3, 5, 14, 30, 0),
	"data_transacao": date(2026, 3, 5),
	"transacao_revisada": 1,
}


def _ler_csv(conteudo: bytes) -> list[list[str]]:
	return list(csv.reader(io.StringIO(conteudo.decode("utf-8-sig")), delimiter=";"))


class TestExtratoCsv(FrappeTestCase):
	def test_cabecalho_e_bom_para_o_excel(self):
		conteudo = gerar_csv_extrato(COLUNAS, [])
		self.assertTrue(conteudo.startswith("﻿".encode()))
		self.assertEqual(
			_ler_csv(conteudo),
			[["Código", "Descrição reduzida", "Valor", "Data/Hora", "Data da transação", "Revisão"]],
		)

	def test_valores_no_formato_brasileiro(self):
		linha = _ler_csv(gerar_csv_extrato(COLUNAS, [TRANSACAO]))[1]
		self.assertEqual(linha[0], "TX-0001")
		self.assertEqual(linha[2], "-1234,50")
		self.assertEqual(linha[3], "05/03/2026 14:30:00")
		self.assertEqual(linha[4], "05/03/2026")
		self.assertEqual(linha[5], "Sim")

	def test_texto_que_parece_formula_vira_literal(self):
		linha = _ler_csv(gerar_csv_extrato(COLUNAS, [TRANSACAO]))[1]
		self.assertEqual(linha[1], "'" + TRANSACAO["descricao_reduzida"])

	def test_campos_vazios_ficam_vazios(self):
		linha = _ler_csv(gerar_csv_extrato(COLUNAS, [{"name": "TX-0002"}]))[1]
		self.assertEqual(linha, ["TX-0002", "", "", "", "", "Não"])


class TestExtratoXlsx(FrappeTestCase):
	def setUp(self):
		self.ws = load_workbook(io.BytesIO(gerar_xlsx_extrato(COLUNAS, [TRANSACAO]))).active

	def test_cabecalho_fixo_e_com_filtro(self):
		self.assertEqual(self.ws.title, "Extrato")
		self.assertEqual(self.ws["A1"].value, "Código")
		self.assertEqual(self.ws["B1"].value, "Descrição reduzida")
		self.assertEqual(self.ws.freeze_panes, "A2")
		self.assertEqual(self.ws.auto_filter.ref, "A1:F2")

	def test_valores_com_tipos_nativos(self):
		self.assertEqual(self.ws["C2"].value, -1234.5)
		self.assertEqual(self.ws["C2"].number_format, "#,##0.00")
		self.assertEqual(self.ws["D2"].value, datetime(2026, 3, 5, 14, 30, 0))
		self.assertEqual(self.ws["E2"].value.date(), date(2026, 3, 5))
		self.assertEqual(self.ws["F2"].value, "Sim")

	def test_texto_que_parece_formula_fica_como_texto(self):
		self.assertEqual(self.ws["B2"].data_type, "s")
		self.assertEqual(self.ws["B2"].value, TRANSACAO["descricao_reduzida"])

	def test_caracteres_de_controle_nao_quebram_a_planilha(self):
		conteudo = gerar_xlsx_extrato(COLUNAS, [{"name": "TX-3", "descricao_reduzida": "Pix\x0bRecebido"}])
		self.assertEqual(load_workbook(io.BytesIO(conteudo)).active["B2"].value, "PixRecebido")


class TestExportarExtratoEndpoint(FrappeTestCase):
	def setUp(self):
		frappe.set_user("Administrator")
		frappe.local.response = frappe._dict()
		self.marcador = f"export-{frappe.generate_hash(length=10)}"

	def tearDown(self):
		frappe.set_user("Administrator")

	def _criar_transacao(self, **campos):
		return (
			frappe.get_doc(
				{
					"doctype": DOCTYPE,
					"id": f"{self.marcador}-{frappe.generate_hash(length=6)}",
					"descricao": "Teste exportação",
					"descricao_reduzida": self.marcador,
					"valor": 10,
					**campos,
				}
			)
			.insert(ignore_permissions=True)
			.name
		)

	def _exportar_csv(self, filtros: dict) -> list[list[str]]:
		exportar_extrato("csv", json.dumps(filtros))
		return _ler_csv(frappe.local.response.filecontent)

	def test_guest_nao_exporta(self):
		frappe.set_user("Guest")
		with self.assertRaises(frappe.PermissionError):
			exportar_extrato("csv")

	def test_formato_invalido(self):
		with self.assertRaises(frappe.ValidationError):
			exportar_extrato("pdf")

	def test_resposta_e_um_download_com_nome_do_formato(self):
		exportar_extrato("XLSX", "{}")
		self.assertEqual(frappe.local.response.type, "download")
		self.assertTrue(frappe.local.response.filename.startswith("extrato-"))
		self.assertTrue(frappe.local.response.filename.endswith(".xlsx"))
		load_workbook(io.BytesIO(frappe.local.response.filecontent))

	def test_exporta_so_o_filtro_e_respeita_as_excluidas(self):
		incluida = self._criar_transacao()
		excluida = self._criar_transacao(excluir_do_total=1)
		self._criar_transacao(descricao_reduzida="outra coisa")

		linhas = self._exportar_csv({"descricao": self.marcador})
		self.assertEqual([linha[0] for linha in linhas[1:]], [incluida])

		linhas = self._exportar_csv({"descricao": self.marcador, "mostrar_excluidas": "1"})
		self.assertEqual({linha[0] for linha in linhas[1:]}, {incluida, excluida})

	def test_exporta_alem_do_tamanho_do_lote_do_grid(self):
		quantidade = EXTRATO_MAX_PAGE_SIZE + 5
		for _ in range(quantidade):
			self._criar_transacao()
		linhas = self._exportar_csv({"descricao": self.marcador})
		self.assertEqual(len(linhas) - 1, quantidade)

	def test_traz_todas_as_colunas_permitidas_mesmo_ocultas(self):
		cabecalho = self._exportar_csv({"descricao": self.marcador})[0]
		self.assertEqual(
			set(cabecalho), {"Código"} | {coluna["label"] for coluna in get_extrato_colunas(True)}
		)
